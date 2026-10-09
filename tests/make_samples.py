"""Vygeneruje vzorové dokumenty s FIKTÍVNYMI osobnými údajmi (platné kontrolné súčty)."""

from __future__ import annotations

import io
import sys
import zipfile
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FONT_CANDIDATES = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf",
                   "C:/Windows/Fonts/segoeui.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
                   "/Library/Fonts/Arial.ttf"]


def find_font() -> str:
    """TrueType font so slovenskou diakritikou (Linux: DejaVu, Windows: Arial, macOS: Arial)."""
    for path in FONT_CANDIDATES:
        if Path(path).exists():
            return path
    raise FileNotFoundError("Nenašiel sa font s diakritikou (DejaVu Sans alebo Arial).")


# ------------------------------------------------------------ platné fiktívne hodnoty
def rodne_cislo(prefix9: str) -> str:
    base = int(prefix9) * 10
    for d in range(10):
        if (base + d) % 11 == 0:
            return f"{prefix9[:6]}/{prefix9[6:]}{d}"
    raise ValueError(prefix9)


def ico(prefix7: str) -> str:
    total = sum(int(c) * w for c, w in zip(prefix7, range(8, 1, -1)))
    return prefix7 + str((11 - total % 11) % 10)


def ic_dph(prefix9: str) -> str:
    for d in range(10):
        if int(prefix9 + str(d)) % 11 == 0:
            return "SK" + prefix9 + str(d)
    raise ValueError(prefix9)


def iban_sk(bank: str, account: str) -> str:
    bban = bank + account.zfill(16)
    check = 98 - int("".join(str(int(c, 36)) for c in bban + "SK00")) % 97
    s = f"SK{check:02d}{bban}"
    return " ".join(s[i:i + 4] for i in range(0, len(s), 4))


RC1 = rodne_cislo("800315123")       # Ján Novák
RC2 = rodne_cislo("856122456")       # Mária Kováčová (žena, mesiac +50)
ICO1 = ico("3575744")
ICDPH1 = ic_dph("202012345")
IBAN1 = iban_sk("1100", "2612345678")

PII = {
    "osoby": ["Ján Novák", "Jána Nováka", "Novákovi", "Mária Kováčová", "Kováčovej", "Peter Horváth",
              "Zuzana Tóthová"],
    "hodnoty": [RC1, RC2, IBAN1, "+421 905 123 456", "0915 987 654", "jan.novak@example.sk",
                "maria.kovacova@example.sk", "EA123456", "BA-123AB", "Hlavná 12", "811 01 Bratislava",
                "Námestie SNP 3/A", "15. 3. 1980", "4111 1111 1111 1111"],
}

CONTRACT = [
    "KÚPNA ZMLUVA č. 2026/114",
    "uzavretá podľa § 588 a nasl. Občianskeho zákonníka",
    "Predávajúci: Ing. Ján Novák, PhD.",
    f"nar. 15. 3. 1980, r. č. {RC1}, č. OP: EA123456",
    "trvalý pobyt: Hlavná 12, 811 01 Bratislava",
    f"IBAN: {IBAN1}, tel.: +421 905 123 456, e-mail: jan.novak@example.sk",
    "Kupujúci: Mária Kováčová",
    f"r. č. {RC2}, bytom Námestie SNP 3/A, 974 01 Banská Bystrica, mobil 0915 987 654",
    "Sprostredkovateľ: ABC Reality s.r.o., Mlynské nivy 5, 821 09 Bratislava",
    f"IČO: {ICO1}, IČ DPH: {ICDPH1}, zastúpená konateľkou Zuzana Tóthová",
    "Čl. I Predmet zmluvy",
    "Predávajúci predáva kupujúcemu osobné motorové vozidlo, EČV BA-123AB, za dohodnutú kúpnu cenu.",
    "Kupujúca zaplatí cenu na účet Jána Nováka do 10 dní. Novákovi zároveň odovzdá podpísaný preberací protokol.",
    "Kľúče od vozidla pán Novák odovzdá Kováčovej pri podpise zmluvy.",
    "V Bratislave dňa 2. 10. 2026",
]


def make_docx(path: Path):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.author = "Ján Novák"
    doc.core_properties.last_modified_by = "Mária Kováčová"
    doc.sections[0].header.paragraphs[0].text = "Advokátska kancelária – spracovala JUDr. Zuzana Tóthová"
    for line in CONTRACT:
        p = doc.add_paragraph()
        if line.startswith("Predávajúci: "):
            # meno rozdelené do viacerých runov s rôznym formátovaním
            p.add_run("Predávajúci: ")
            p.add_run("Ing. Ján ").bold = True
            r = p.add_run("Novák")
            r.bold = True
            r.font.size = Pt(13)
            p.add_run(", PhD.")
        else:
            p.add_run(line)
    t = doc.add_table(rows=3, cols=2)
    rows = [("Meno a priezvisko", "Peter Horváth"), ("Telefón", "+421 905 123 456"),
            ("Kontakt", "maria.kovacova@example.sk")]
    for row, (a, b) in zip(t.rows, rows):
        row.cells[0].text, row.cells[1].text = a, b
    doc.add_paragraph("S pozdravom")
    doc.add_paragraph("MUDr. Peter Horváth")
    buf = io.BytesIO()
    doc.save(buf)
    # sledovaná zmena: vymazaný text s osobným údajom
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                s = data.decode("utf-8")
                tracked = ('<w:p><w:r><w:t xml:space="preserve">Pôvodný kupujúci: </w:t></w:r>'
                           '<w:del w:id="91" w:author="Ján Novák" w:date="2026-10-01T10:00:00Z"><w:r>'
                           '<w:delText>Jozef Malý, tel. 0903 111 222</w:delText></w:r></w:del></w:p>')
                s = s.replace("<w:sectPr", tracked + "<w:sectPr", 1)
                data = s.encode("utf-8")
            zout.writestr(info, data)
    path.write_bytes(out.getvalue())


def make_pdf(path: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("DejaVu", find_font()))
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setAuthor("Ján Novák")
    y = 800
    for line in CONTRACT:
        size = 11
        c.setFont("DejaVu", size)
        for chunk in _wrap(line, 85):
            c.drawString(50, y, chunk)
            y -= 18
        if y < 80:
            c.showPage()
            y = 800
    c.save()


def _wrap(text, width):
    words, line, out = text.split(), "", []
    for w in words:
        if len(line) + len(w) + 1 > width:
            out.append(line)
            line = w
        else:
            line = (line + " " + w).strip()
    out.append(line)
    return out


def make_scan(src_pdf: Path, path: Path):
    from pdf2image import convert_from_path
    from PIL import ImageFilter

    sys.path.insert(0, str(ROOT.parent))
    from anonymizer_sk.runtime import configure

    poppler = configure()          # na Windows: Poppler z build/tools alebo ANONYMIZER_POPPLER

    pages = [im.convert("L").rotate(0.4, fillcolor=255, expand=False).filter(ImageFilter.GaussianBlur(0.4))
             for im in convert_from_path(str(src_pdf), dpi=200, **poppler)]
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=200)


def make_eml(path: Path, docx_path: Path):
    msg = EmailMessage()
    msg["From"] = "Ján Novák <jan.novak@example.sk>"
    msg["To"] = "Mária Kováčová <maria.kovacova@example.sk>"
    msg["Cc"] = "Peter Horváth <peter.horvath@example.sk>"
    msg["Subject"] = "Zmluva pre pani Kováčovú – auto BA-123AB"
    msg["Received"] = "from mail.example.sk (mail.example.sk [203.0.113.45]) by mx.example.sk"
    msg["X-Originating-IP"] = "[203.0.113.45]"
    msg["Message-ID"] = "<abc123@example.sk>"
    body = (f"Dobrý deň pani Kováčová,\n\nposielam návrh zmluvy. Moje rodné číslo je {RC1} a platbu "
            f"prosím pošlite na IBAN {IBAN1}.\nV prípade otázok volajte na 0905 123 456.\n\n"
            "S pozdravom\nJán Novák\n")
    msg.set_content(body)
    msg.add_alternative(
        "<html><body><p>Dobrý deň pani <b>Kováčová</b>,</p><p>posielam návrh zmluvy. Moje rodné číslo je "
        f"{RC1} a platbu prosím pošlite na IBAN {IBAN1}.</p><p>V prípade otázok volajte na 0905 123 456 alebo "
        "píšte na <a href=\"mailto:jan.novak@example.sk\">jan.novak@example.sk</a>.</p>"
        "<p>S pozdravom<br>\nJán Novák</p></body></html>", subtype="html")
    msg.add_attachment(docx_path.read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                       filename="zmluva_Novak.docx")
    msg.add_attachment(b"\x89PNG\r\n\x1a\nfake", maintype="image", subtype="png", filename="op_scan.png")
    path.write_bytes(bytes(msg))


def main(out_dir: Path):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "poznamka.txt").write_text(
        "Poznámka z telefonátu: pán Ján Novák (tel. +421 905 123 456) žiada zmenu termínu.\n"
        f"Jeho rodné číslo {RC1} treba overiť. Kontakt na manželku: Jana Nováková, jana.novakova@example.sk\n"
        "Stretnutie s Novákom a Novákovou bude v piatok.\n", encoding="utf-8")
    make_docx(out_dir / "zmluva.docx")
    make_pdf(out_dir / "zmluva_text.pdf")
    make_scan(out_dir / "zmluva_text.pdf", out_dir / "zmluva_sken.pdf")
    make_eml(out_dir / "email.eml", out_dir / "zmluva.docx")
    print("Vzorky vytvorené v", out_dir)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "samples")
